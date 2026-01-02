from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_word_doc():
    doc = Document()
    
    # Title
    title = doc.add_heading('SAE Aero Design 2026: Advanced Class Spot Check Questions', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Intro/Disclaimer
    intro = doc.add_paragraph(
        "These questions are based on the 2026 SAE Aero Design Rules (Version 2026.0). "
        "Page references correspond to the page numbers in the official PDF ruleset."
    )
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph() # Spacer

    # Data Structure: Category -> List of (Question, Answer, Page Number)
    content = {
        "1. Controls (Autopilot, Radio, Servos)": [
            ("For your autonomous system's manual override, describe the specific physical configuration of the switch on your transmitter.", 
             "It must be a red-colored momentary switch, and it must be configured such that the pilot holds the switch to keep the aircraft in autonomous flight mode (releasing it returns to manual).", "37"),
            ("What is the minimum number of degrees of freedom your active navigation system must control to be considered 'autonomous' under the rules?", 
             "It must control at least two (2) degrees of freedom.", "37"),
            ("Demonstrate your radio fail-safe. What specific throttle action must occur immediately upon signal loss?", 
             "The system must reduce the throttle to zero immediately.", "11"),
            ("If you are using an Autopilot system, how must it be powered and armed?", 
             "It must use a discrete and removable Red arming plug subject to Section 2.22 requirements.", "38"),
            ("You have a manual reset switch for your autopilot on the aircraft. Where must it be located?", 
             "It must be located externally in accordance with Section 2.24 (clearly visible, 9 inches from prop).", "38, 14"),
            ("What frequency band is prohibited for your Autopilot data link?", 
             "It shall not use the same 2.4 GHz band as the pilot transmitter.", "38"),
            ("Show me the mechanical retention devices on your control clevises. What are they called and why are they there?", 
             "They are Clevis Keepers required to prevent accidental opening of the clevis during flight.", "12"),
            ("Does your aircraft possess a ground steering mechanism, and is it allowed to rely solely on the rudder?", 
             "Yes, it must have a ground steering mechanism if it has wheels; it cannot rely solely on aerodynamic control surfaces (rudder) for steering.", "11"),
            ("Your team decides to use an FPV system. What are the frequency constraints regarding the flight control system?", 
             "FPV systems cannot use the same frequency as the flight control system, and the use of 2.4 GHz for FPV video is prohibited.", "38"),
            ("How did you determine that your servos are strong enough for this aircraft?", 
             "We performed analysis and/or testing described in the Design Report to demonstrate they are sized for expected aerodynamic flight loads.", "12")
        ],
        "2. Structures (Airframe, Dimensions, Materials)": [
            ("What is the maximum planform wingspan allowed for your Advanced Class aircraft?", 
             "Less than 120 inches.", "34"),
            ("According to the Section 8 Design Requirements, what is the maximum weight limit for your aircraft?", 
             "The Advanced Class aircraft maximum weight shall not exceed 3.50 lbs.", "34"),
            ("Are you using rubber bands to hold your wing onto the fuselage?", 
             "No, elastic material such as rubber bands shall not retain the wing to the fuselage.", "34"),
            ("Show me the Center of Gravity (CG) markings on your fuselage. What are the required dimensions for these marks?", 
             "They must be a minimum of 0.5 inches in diameter.", "10"),
            ("What is the tolerance allowed for the placement of your CG marking compared to the empty CG position on your 2D drawings?", 
             "It must be centered at the Empty CG position +/- 0.25 inches.", "10"),
            ("We need to verify your Empty CG. Does the aircraft need to be flyable at this exact Empty CG location?", 
             "Yes, all aircraft shall be flyable at their designated Empty CG position.", "10"),
            ("Did you use any metal propellers in this design?", 
             "No, metal propellers are prohibited.", "12"),
            ("Is lead (Pb) used anywhere in this aircraft, perhaps for ballast in the nose?", 
             "No, the use of lead in any portion of the aircraft (including payload) is prohibited.", "12"),
            ("If your aircraft is damaged and you repair it, can you make design changes to improve it?", 
             "Repairs must not deviate from the baseline design; major repairs must undergo safety inspection.", "14"),
            ("Where is the university name displayed on the aircraft?", 
             "It must be clearly displayed externally on the wings or fuselage (or unique University initials).", "10")
        ],
        "3. Payload (Design, Loading, DLZ)": [
            ("What is the maximum linear dimension allowed for any single payload unit?", 
             "Twelve (12) inches or less.", "36"),
            ("What is the time limit for unloading your payload during the demonstration?", 
             "Under one (1) minute.", "36"),
            ("Your payload contains electronics. Are these allowed, and is there a restriction on how they operate?", 
             "Electronics are allowed, but the payload shall not be manually operated.", "36"),
            ("We are setting up your Designated Landing Zone (DLZ). How many stakes are required, and what is the minimum length of each stake?", 
             "A minimum of nine (9) stakes are required, and they must be eight (8) inches or longer.", "37"),
            ("Can you use Velcro or magnets on your DLZ surface to help catch the payload?", 
             "No, no nets or modified surfaces that can adhere to or catch the payload (Velcro, magnets, tape) are allowed.", "37"),
            ("How many payloads can your aircraft carry during a single mission segment?", 
             "Only one (1) payload shall be carried at a time.", "36"),
            ("Explain the font size requirements for marking your payloads.", 
             "Team number and unique payload number must be in minimum one (1) inch font.", "36"),
            ("If you successfully Release a payload but it rolls off the DLZ, does it score?", 
             "No, the payload must remain entirely within the DLZ; if it touches the ground outside, it is disqualified from mission scoring.", "36-37"),
            ("Can your DLZ contain any batteries or electronics to help with the 'Capture' phase?", 
             "No, no electronics, batteries, or magnets are allowed as part of the installed DLZ.", "37"),
            ("When is the payload weighed during the competition flow?", 
             "Prior to flight (if present during takeoff) and after the flight attempt (if present during landing).", "36")
        ],
        "4. Electronics/Power Management": [
            ("What are the exact voltage and capacity limits for your Advanced Class propulsion battery?", 
             "4 cell (14.8 volt) Lithium-Polymer, maximum capacity of 3000 mAh.", "34"),
            ("Describe the required location of your Red Arming Plug relative to the propeller.", 
             "It must be a minimum of nine (9) inches away from any propeller at any point in its rotational plane.", "13"),
            ("Can you disconnect the wiring harness to arm/disarm the system instead of using a plug?", 
             "No, disconnecting wiring harnesses to arm/disarm is prohibited.", "13"),
            ("On which wire must the arming plug be integrated?", 
             "On the positive (RED) wire between the battery and the ESC.", "13"),
            ("How is your receiver system powered? If you use a separate battery, what are the specs?", 
             "A separate battery OR separate BEC is required. If a battery is used, it must be min 1000 mAh LiPo or LiFe.", "14"),
            ("Where is the On/Off switch for the receiver located?", 
             "Mounted to the aircraft exterior, at least nine (9) inches from any propeller.", "14"),
            ("Are you using a Power Limiter?", 
             "Advanced Class is not required to fly with a power limiter.", "34"),
            ("Inspecting the arming plug interface: how many male leads can the non-removable portion have?", 
             "It shall not have more than one male lead.", "13"),
            ("Are your batteries 'positively secured'? What does that mean according to the rules?", 
             "They must be unable to move under all flight loads.", "12"),
            ("Is the battery bay free of protrusions?", 
             "Yes, it must be free of hardware that could penetrate the battery in a crash.", "12")
        ],
        "5. Propulsion (Motors, Props)": [
            ("What is the maximum number of motors and propellers allowed on your Advanced Class aircraft?", 
             "Maximum of three (3) motors and three (3) propellers.", "34"),
            ("Can you use a gearbox? If so, does the prop have to spin at motor RPM?", 
             "Gearboxes are allowed. (Unlike Regular Class, Advanced Class does not explicitly require 1:1 RPM).", "34"),
            ("What is the maximum thrust line angle allowed relative to the horizontal during conventional takeoff?", 
             "Up to ten (10) degrees from horizontal.", "35"),
            ("You have a VTOL-style lifting prop. Can it be active during the conventional takeoff run?", 
             "No, any propulsion device oriented greater than 10 degrees from horizontal shall not be active during Conventional Takeoff through liftoff.", "35"),
            ("What safety device is required on the propeller shaft?", 
             "A spinner or a rounded model aircraft type safety nut (Nylon-insert lock-nuts are prohibited).", "11"),
            ("Are you allowed to perform a motor run-up in the On-Deck area?", 
             "No, motor runup and testing shall not be allowed in On-Deck.", "18"),
            ("Can you use a metal propeller?", 
             "No.", "12"),
            ("Does your throttle setting go to zero if the radio signal is lost?", 
             "Yes, fail-safe must reduce throttle to zero.", "11"),
            ("Can you use a 'pusher' configuration?", 
             "Yes, provided all safety distance rules (9 inches for arming plugs/switches) and arming plug visibility rules are met.", "13-14"),
            ("Are you using any form of stored energy for propulsion other than the batteries (e.g., capacitors, pressure vessels)?", 
             "No, stored energy restriction prohibits this.", "12")
        ],
        "6. Other (Mission, General Rules, Documentation)": [
            ("You must show us a video before you can fly. What specifically must the 'Proof of Flight' video demonstrate?", 
             "1. Conventional Takeoff with sustained stable flight for 10 seconds. 2. Return to Base with controlled landing without damage.", "34"),
            ("What is the maximum duration allowed for your Proof of Flight video?", 
             "No more than 1.5 minutes.", "34"),
            ("If you fail a 'spot check' on a general requirement during inspection, what is the penalty?", 
             "There shall be a point penalty for each item failed, and the aircraft must be brought into compliance.", "30"),
            ("How many team members are allowed to go with the pilot to the runway for an Advanced Class flight?", 
             "One (1) team member (Escort) plus the Pilot.", "17"),
            ("If you enter a No-Fly Zone (NFZ) for the first time, what is the penalty?", 
             "Disqualified flight attempt and zero points for that flight.", "20"),
            ("If you enter the NFZ a second time?", 
             "Disqualification from the entire event and loss of all points.", "20"),
            ("Are you allowed to use a laser to help aim the aircraft for the drop?", 
             "No, use of lasers for marking/highlighting landing zones or directing aircraft is prohibited.", "12"),
            ("Do you have your FDRR (Flight Demonstration Readiness Review) presentation ready? What happens if you exceed the time limit by 1 minute?", 
             "We will be assessed a five (5) point penalty and stopped after one additional minute.", "28"),
            ("Who is allowed to act as the Payload Operator (PO)?", 
             "One (1) team member; they cannot rely on line-of-sight view to the aircraft or DLZ.", "38"),
            ("If your aircraft crashes and you repair it, do you need to get inspected again?", 
             "Yes, all major repairs shall undergo safety inspection before the aircraft is cleared for flight.", "14")
        ]
    }

    # Style application
    for category, items in content.items():
        doc.add_heading(category, level=1)
        for i, (question, answer, page) in enumerate(items, 1):
            p = doc.add_paragraph()
            # Question in bold
            runner_q = p.add_run(f"Q{i}: {question}")
            runner_q.bold = True
            
            # Line break
            p.add_run("\n")
            
            # Answer
            p.add_run(f"A: {answer} ")
            
            # Page Citation in Italics/Gray
            runner_cite = p.add_run(f"[Page {page}]")
            runner_cite.italic = True
            runner_cite.font.color.rgb = RGBColor(100, 100, 100)
            
            # Spacing between Q&As
            p.paragraph_format.space_after = Pt(12)

    doc.save('SAE_Aero_Design_Advanced_Spot_Check_Questions.docx')

if __name__ == "__main__":
    create_word_doc()